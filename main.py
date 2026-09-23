import fitdecode
import zipfile

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

mapje = Path(__file__).parent
downloads = Path.home() / "Downloads"

excel_bestand = mapje / "garmin_activiteiten.xlsx"

backup_map = mapje / "FIT_backup"
backup_map.mkdir(exist_ok=True)

if excel_bestand.exists():

    wb = load_workbook(excel_bestand)

    ws = wb["Activiteiten"]

else:

    wb = Workbook()

    ws = wb.active
    ws.title = "Activiteiten"

    koppen = [
        "Datum",
        "Sport",
        "Duur/tijd",
        "Afstand (km)",
        "Tempo / snelheid",
        "Gem. hartslag",
        "Max hartslag",
        "Aeroob effect",
        "Anaeroob effect",
        "RPE",
        "Gevoel",
        "Activiteit ID",
        "Weeknummer",
        "Load"
    ]

    ws.append(koppen)

bestaande_ids = {
    ws.cell(rij, 12).value
    for rij in range(2, ws.max_row + 1)
}

def seconden_naar_tijd(seconden):

    if seconden is None:
        return None

    minuten = int(seconden // 60)

    seconden_rest = int(
        round(seconden % 60)
    )

    if seconden_rest == 60:

        minuten += 1
        seconden_rest = 0

    return f"{minuten:02d}:{seconden_rest:02d}"

zip_bestanden = list(
    downloads.glob("*.zip")
)

print(
    "Gevonden ZIP-bestanden:",
    len(zip_bestanden)
)

for zip_bestand in zip_bestanden:

    print("\n-----------------------------------")
    print(
        "ZIP verwerken:",
        zip_bestand.name
    )
    print("-----------------------------------")

    succesvol_verwerkt = True

    try:

        with zipfile.ZipFile(
            zip_bestand,
            "r"
        ) as zip_file:

            fit_bestanden = [
                naam
                for naam in zip_file.namelist()
                if naam.lower().endswith(".fit")
            ]

            print(
                "FIT-bestanden gevonden:",
                len(fit_bestanden)
            )

            for fit_bestand in fit_bestanden:

                activiteit_id = Path(
                    fit_bestand
                ).name

                print(
                    "\nFIT:",
                    activiteit_id
                )

                if activiteit_id in bestaande_ids:

                    print(
                        "Activiteit bestaat al in Excel en wordt overgeslagen:",
                        activiteit_id
                    )

                    continue

                backup_bestand = (
                    backup_map / activiteit_id
                )

                if backup_bestand.exists():

                    print(
                        "FIT bestaat al in backup:",
                        activiteit_id
                    )

                else:

                    with zip_file.open(
                        fit_bestand
                    ) as bron:

                        with open(
                            backup_bestand,
                            "wb"
                        ) as doel:

                            doel.write(
                                bron.read()
                            )

                    print(
                        "FIT opgeslagen in backup:",
                        activiteit_id
                    )

                activiteit = {}

                with fitdecode.FitReader(
                    str(backup_bestand)
                ) as fit:

                    for frame in fit:

                        if isinstance(
                            frame,
                            fitdecode.FitDataMessage
                        ):

                            if frame.name == "session":

                                for field in frame.fields:

                                    if field.value is not None:

                                        activiteit[
                                            field.name
                                        ] = field.value

                datum = activiteit.get(
                    "start_time"
                )

                if datum is not None:

                    datum = datum.date()

                sport = activiteit.get(
                    "sport"
                )

                sport_string = str(
                    sport
                ).lower()

                sport_waarde = getattr(
                    sport,
                    "value",
                    None
                )

                if (
                    sport_waarde == 17
                    or "hiking" in sport_string
                ):

                    sport_naam = "Wandelen"

                elif (
                    sport_waarde == 11
                    or "walking" in sport_string
                ):

                    sport_naam = "Wandelen"

                elif "running" in sport_string:

                    sport_naam = "Hardlopen"

                elif "cycling" in sport_string:

                    sport_naam = "Fietsen"

                elif "swimming" in sport_string:

                    sport_naam = "Zwemmen"

                elif (
                    "training" in sport_string
                    or "fitness" in sport_string
                ):

                    sport_naam = "Fitness"

                else:

                    sport_naam = "Overig"


                print(
                    "Garmin sport:",
                    repr(sport)
                )

                print(
                    "Excel sport:",
                    sport_naam
                )

                duur_seconden = activiteit.get(
                    "total_timer_time"
                )

                if duur_seconden is not None:

                    duur = (
                        duur_seconden / 86400
                    )

                else:

                    duur = None

                total_distance = activiteit.get(
                    "total_distance"
                )

                if (
                    sport_string in [
                        "training",
                        "fitness"
                    ]
                    or sport_naam == "Badminton"
                ):

                    afstand_km = None

                elif total_distance is not None:

                    afstand_km = (
                        total_distance / 1000
                    )

                else:

                    afstand_km = None

                snelheid = activiteit.get(
                    "enhanced_avg_speed"
                )

                if snelheid is None:

                    snelheid = activiteit.get(
                        "avg_speed"
                    )

                tempo_snelheid = ""

                if (
                    snelheid is not None
                    and snelheid > 0
                ):


                    # HARDLOPEN

                    if sport_string == "running":

                        seconden_per_km = (
                            1000 / snelheid
                        )

                        tempo_string = (
                            seconden_naar_tijd(
                                seconden_per_km
                            )
                        )

                        tempo_snelheid = (
                            tempo_string
                            + " min/km"
                        )


                    # FIETSEN

                    elif sport_string == "cycling":

                        km_per_uur = (
                            snelheid * 3.6
                        )

                        tempo_snelheid = (
                            f"{km_per_uur:.1f} km/u"
                        )


                    # ZWEMMEN

                    elif sport_string == "swimming":

                        seconden_per_100m = (
                            100 / snelheid
                        )

                        tempo_string = (
                            seconden_naar_tijd(
                                seconden_per_100m
                            )
                        )

                        tempo_snelheid = (
                            tempo_string
                            + " min/100m"
                        )


                    # WANDELEN / HIKING /
                    # FITNESS / BADMINTON

                    else:

                        tempo_snelheid = ""

                gem_hartslag = activiteit.get(
                    "avg_heart_rate"
                )

                max_hartslag = activiteit.get(
                    "max_heart_rate"
                )

                aeroob_effect = activiteit.get(
                    "total_training_effect"
                )

                anaeroob_effect = activiteit.get(
                    "total_anaerobic_training_effect"
                )

                workout_rpe = activiteit.get(
                    "workout_rpe"
                )

                if workout_rpe is not None:

                    rpe = workout_rpe / 10

                else:

                    rpe = None

                workout_feel = activiteit.get(
                    "workout_feel"
                )

                gevoelens = {

                    0: "Zeer zwak",
                    25: "Zwak",
                    50: "Normaal",
                    75: "Sterk",
                    100: "Zeer sterk"

                }

                if workout_feel is not None:

                    gevoel = gevoelens.get(
                        workout_feel,
                        workout_feel
                    )

                else:

                    gevoel = ""

                ws.append([

                    datum,
                    sport_naam,
                    duur,
                    afstand_km,
                    tempo_snelheid,
                    gem_hartslag,
                    max_hartslag,
                    aeroob_effect,
                    anaeroob_effect,
                    rpe,
                    gevoel,
                    activiteit_id

                ])

                nieuwe_rij = ws.max_row

                ws.cell(
                    nieuwe_rij,
                    13
                ).value = (
                    f"=WEEKNUM(A{nieuwe_rij})"
                )

                ws.cell(
                    nieuwe_rij,
                    14
                ).value = (
                    f'=IF(J{nieuwe_rij}="","",C{nieuwe_rij}*J{nieuwe_rij}*24)'
                )

                bestaande_ids.add(
                    activiteit_id
                )


                print(
                    "Activiteit toegevoegd aan Excel:",
                    activiteit_id
                )

        if succesvol_verwerkt:

            zip_bestand.unlink()

            print(
                "ZIP succesvol verwerkt en verwijderd:",
                zip_bestand.name
            )


    except Exception as fout:

        succesvol_verwerkt = False

        print(
            "\nFOUT bij verwerken van:",
            zip_bestand.name
        )

        print(
            fout
        )

        print(
            "ZIP wordt NIET verwijderd."
        )

for rij in range(
    2,
    ws.max_row + 1
):

    ws.cell(
        rij,
        1
    ).number_format = "dd-mm-yyyy"

    ws.cell(
        rij,
        3
    ).number_format = "[h]:mm:ss"

if "ActiviteitenTabel" in ws.tables:

    tabel = ws.tables[
        "ActiviteitenTabel"
    ]

    print(
        "Bestaande tabel gevonden:",
        "ActiviteitenTabel"
    )

else:

    print(
        "ActiviteitenTabel niet gevonden."
    )

    print(
        "Er wordt een nieuwe tabel aangemaakt."
    )

    tabel = Table(
        displayName="ActiviteitenTabel",
        ref=f"A1:N{ws.max_row}"
    )

    stijl = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False
    )

    tabel.tableStyleInfo = stijl

    ws.add_table(tabel)

tabel.ref = f"A1:N{ws.max_row}"

print(
    "Tabelbereik:",
    tabel.ref
)


try:

    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"

except Exception:

    pass

wb.save(
    excel_bestand
)

print(
    "\n-----------------------------------"
)

print(
    "Klaar!"
)

print(
    "Excel opgeslagen:",
    excel_bestand
)
